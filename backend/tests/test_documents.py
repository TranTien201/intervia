import io

import docx
from pypdf import PdfWriter

from app.documents import MIN_CHARS_PER_PAGE, read_text, squeeze
from app.domain.files import DOCX, PDF
from tests.samples import LINE, make_docx, make_pdf


def test_pdf_with_text_is_readable():
    result = read_text(make_pdf([LINE, LINE]), PDF)
    assert "payment APIs with Python" in result.text
    assert result.page_count == 1
    assert result.looks_readable


def test_pdf_without_text_is_not_readable():
    """A scanned CV: pages, but no text layer."""
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    writer.add_blank_page(width=595, height=842)
    buffer = io.BytesIO()
    writer.write(buffer)
    result = read_text(buffer.getvalue(), PDF)
    assert result.text == ""
    assert result.page_count == 2
    assert not result.looks_readable


def test_corrupt_file_is_reported_not_raised():
    result = read_text(b"%PDF-1.7 not really a pdf", PDF)
    assert not result.looks_readable
    assert read_text(b"PK\x03\x04broken", DOCX).looks_readable is False


def test_docx_keeps_tables_in_order_and_merged_cells_once():
    data = make_docx([LINE, LINE], table=[["Python", "5 năm"], ["Go", "2 năm"]])
    document = docx.Document(io.BytesIO(data))
    merged = document.tables[0]
    merged.cell(1, 0).merge(merged.cell(1, 1))  # "Go" + "2 năm" become one cell
    buffer = io.BytesIO()
    document.save(buffer)

    result = read_text(buffer.getvalue(), DOCX)
    assert result.text.endswith("Python | 5 năm\nGo\n2 năm")
    assert result.looks_readable


def test_too_little_text_is_not_readable():
    assert not read_text(make_docx(["x" * (MIN_CHARS_PER_PAGE - 1)]), DOCX).looks_readable


def test_broken_font_characters_are_not_readable():
    """Private-use glyphs are what PDFs with a broken font encoding produce."""
    garbled = "".join(chr(0xE000 + i % 50) if i % 10 == 0 else "a" for i in range(500))
    assert not read_text(make_docx([garbled]), DOCX).looks_readable


def test_squeeze():
    assert squeeze("  a \t b\xa0c \n\n\n\n d  \n") == "a b c\n\nd"
    # Decomposed Vietnamese (e + combining marks) becomes the composed form.
    assert squeeze("nghiệm") == "nghiệm"
